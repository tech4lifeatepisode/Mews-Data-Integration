from __future__ import annotations

import json
import os
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openai import OpenAI


BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / ".env"
WORKBOOK_PATH = BASE_DIR / "Company ID.xlsx"
BATCH_SIZE = 180


def normalize(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def main() -> None:
    load_dotenv(ENV_PATH)
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    model = os.getenv("OPENAI_MODEL", "gpt-5.2-2025-12-11").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY missing in .env")

    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb.active

    ids: list[str] = []
    rows: list[int] = []
    for r in range(2, ws.max_row + 1):
        v = normalize(ws.cell(r, 1).value)
        if v:
            ids.append(v)
            rows.append(r)

    client = OpenAI(api_key=api_key)
    answers: dict[str, str] = {}

    system = (
        "You classify if each identifier is a company tax ID (legal entity) in Spain.\n"
        "Return JSON only: {\"results\":[{\"id\":\"...\",\"answer\":\"Yes|No\"}, ...]}.\n"
        "Use 'Yes' only when the ID pattern strongly indicates company/legal-entity identifier.\n"
        "Use 'No' for personal IDs, passports, NIE/NIF for individuals, or unknown/ambiguous values."
    )

    unique_ids = sorted(set(ids))
    for i in range(0, len(unique_ids), BATCH_SIZE):
        chunk = unique_ids[i : i + BATCH_SIZE]
        payload = json.dumps({"ids": chunk}, ensure_ascii=False)
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": payload},
            ],
        )
        content = (resp.choices[0].message.content or "{}").strip()
        data = json.loads(content)
        items = data.get("results", [])
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            k = normalize(item.get("id"))
            a = normalize(item.get("answer")).capitalize()
            if k:
                answers[k] = "Yes" if a == "Yes" else "No"

    ws.cell(1, 2, "Company ID")
    for id_value, row in zip(ids, rows):
        ws.cell(row, 2, answers.get(id_value, "No"))

    wb.save(WORKBOOK_PATH)
    wb.close()
    print(f"Updated {WORKBOOK_PATH} with {len(rows)} rows.")


if __name__ == "__main__":
    main()
