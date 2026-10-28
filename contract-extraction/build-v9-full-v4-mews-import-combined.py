"""
V9 Mews xlsx from V4 CSV exports — Mews-import-safe (clears Source/Segment, fixes doc type, categories).

Outputs:
  - V9_Full_* — all rows, sanitized (may still fail historical check-in if arrival is too far in past)
  - V9_Importable_* — arrival >= 2026-10-21 only (matches V8 sample window; recommended for Mews import)

  python contract-extraction/build-v9-full-v4-mews-import-combined.py
"""
from __future__ import annotations

import csv
import importlib.util
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

_v8_path = Path(__file__).resolve().parent / "build-v8-random-10-mews-import-sample.py"
_spec = importlib.util.spec_from_file_location("v8_sample", _v8_path)
_v8 = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_v8)
PREPAYMENT_TYPE = _v8.PREPAYMENT_TYPE
mews_import_ready = _v8.mews_import_ready

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Mews_import_v2.10.0.xlsx"
CUSTOMERS_CSV = ROOT / "contract-extraction" / "v4-customers-export-21-10-2026.csv"
COMPANIES_CSV = ROOT / "contract-extraction" / "v4-companies-export-21-10-2026.csv"
MIN_ARRIVAL_DATE = datetime(2026, 10, 21)

OUTPUT_FULL = Path.home() / "Downloads" / "V9_Full_Mews_import_V4_B2C_and_B2B_21_10_2026_combined.xlsx"
OUTPUT_IMPORTABLE = Path.home() / "Downloads" / "V9_Importable_Mews_import_V4_B2C_and_B2B_21_10_2026_combined.xlsx"
OUTPUT_REPO_FULL = ROOT / "Mews_import_V9_Full_V4_B2C_B2B_21_10_2026_combined.xlsx"
OUTPUT_REPO_IMPORTABLE = ROOT / "Mews_import_V9_Importable_V4_B2C_B2B_21_10_2026_combined.xlsx"

CATEGORY_ALIASES = {
    "2 bedroom (double / single)": "The Two Bedroom (Double Single)",
    "2 bedroom (double/single)": "The Two Bedroom (Double Single)",
    "2 bedroom (2+1)": "The Two Bedroom (Double Single)",
    "the two bedroom (double single)": "The Two Bedroom (Double Single)",
}

INVALID_DOC_TYPES = {"taxidentificationnumber", "tax identification number"}


def parse_date(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return None


def canonicalize_category(raw) -> str | None:
    if raw is None or str(raw).strip() == "":
        return None
    s = str(raw).strip()
    key = s.lower()
    if key in CATEGORY_ALIASES:
        return CATEGORY_ALIASES[key]
    for alias, canon in CATEGORY_ALIASES.items():
        if alias in key:
            return canon
    return s


def sanitize_mews_row(row: list, idx: dict[str, int]) -> None:
    mews_import_ready(row, idx)

    for key in ("Source", "Business Segment", "Channel manager integration ID", "Channel manager number", "Channel number"):
        if key in idx:
            row[idx[key]] = None

    if "Document Type" in idx:
        dt = row[idx["Document Type"]]
        if dt and str(dt).strip().lower().replace(" ", "") in INVALID_DOC_TYPES:
            doc = row[idx.get("Document Number", -1)] if "Document Number" in idx else ""
            doc_s = str(doc or "").strip()
            row[idx["Document Type"]] = "Passport" if doc_s and doc_s[0].isalpha() and len(doc_s) >= 6 else "IdentityCard"

    if "Included Products" in idx:
        row[idx["Included Products"]] = None

    if "Allocated Room Category(*)" in idx and row[idx["Allocated Room Category(*)"]]:
        canon = canonicalize_category(row[idx["Allocated Room Category(*)"]])
        if canon:
            row[idx["Allocated Room Category(*)"]] = canon


def load_customer_headers(ws) -> list[str]:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def read_customer_csv(path: Path, headers: list[str]) -> list[list]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [[rec.get(h) or None for h in headers] for rec in reader]


def read_company_csv(path: Path, headers: list[str]) -> list[list]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [[rec.get(h) or None for h in headers] for rec in reader]


def transform_customers(rows: list[list], headers: list[str]) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    out: list[tuple] = []
    for row in rows:
        r = list(row)
        sanitize_mews_row(r, idx)
        out.append(tuple(r))
    return out


def filter_importable(customers: list[tuple], headers: list[str]) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    arr_i = idx["Arrival Date(*)"]
    kept = []
    for row in customers:
        a = parse_date(row[arr_i])
        if a and a >= MIN_ARRIVAL_DATE:
            kept.append(row)
    return kept


def write_sheet(ws, headers: list[str], data_rows: list[tuple]) -> None:
    for c, h in enumerate(headers, start=1):
        if ws.cell(row=1, column=c).value != h:
            raise SystemExit(f"Header mismatch col {c}: {ws.cell(row=1, column=c).value!r} != {h!r}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    room_col = headers.index("Allocated Room Number") + 1 if "Allocated Room Number" in headers else None
    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if room_col == c_idx and val is not None and str(val).strip().startswith("-"):
                cell.number_format = "@"
                cell.value = str(val).strip()


def write_workbook(path: Path, cust_headers: list[str], customers: list[tuple], co_headers: list[str], companies: list[tuple]) -> None:
    shutil.copy2(TEMPLATE, path)
    wb = openpyxl.load_workbook(path)
    write_sheet(wb["Customers"], cust_headers, customers)
    write_sheet(wb["Companies"], co_headers, companies)
    wb.save(path)


def main() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")
    if not CUSTOMERS_CSV.is_file():
        raise SystemExit(f"Run upload-v4 first. Missing: {CUSTOMERS_CSV}")
    if not COMPANIES_CSV.is_file():
        raise SystemExit(f"Missing: {COMPANIES_CSV}")

    tpl_wb = openpyxl.load_workbook(TEMPLATE, read_only=True)
    cust_headers = load_customer_headers(tpl_wb["Customers"])
    co_headers = load_customer_headers(tpl_wb["Companies"])
    tpl_wb.close()

    raw_customers = read_customer_csv(CUSTOMERS_CSV, cust_headers)
    company_rows = [tuple(r) for r in read_company_csv(COMPANIES_CSV, co_headers)]
    customers_all = transform_customers(raw_customers, cust_headers)
    customers_importable = filter_importable(customers_all, cust_headers)

    for path in (OUTPUT_FULL, OUTPUT_REPO_FULL):
        write_workbook(path, cust_headers, customers_all, co_headers, company_rows)
    for path in (OUTPUT_IMPORTABLE, OUTPUT_REPO_IMPORTABLE):
        write_workbook(path, cust_headers, customers_importable, co_headers, company_rows)

    idx = {h: i for i, h in enumerate(cust_headers)}
    prep = sum(
        1
        for r in customers_all
        if idx.get("Prepayment Type") is not None and r[idx["Prepayment Type"]] == PREPAYMENT_TYPE
    )
    print(f"Wrote {OUTPUT_FULL} ({len(customers_all)} customers — sanitized; past arrivals may fail Mews historical window)")
    print(f"Wrote {OUTPUT_IMPORTABLE} ({len(customers_importable)} customers — arrival >= {MIN_ARRIVAL_DATE.date()}, use for import)")
    print(f"  Repo copies: {OUTPUT_REPO_FULL.name}, {OUTPUT_REPO_IMPORTABLE.name}")
    print(f"  Companies: {len(company_rows)}")
    print(f"  Prepayment Type Deposit: {prep}/{len(customers_all)}")


if __name__ == "__main__":
    main()
