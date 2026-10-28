"""
V17 Mews import workbook from live V6 (607 customers; arrivals normalized to 2026-11-01 where V6 rule applied).

  node contract-extraction/export-v6-tables-csv-21-10-2026.mjs V17
  python contract-extraction/build-V17-mews-import-from-v6.py
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

from _mews_v10_from_v4_common import (
    companies_for_customers,
    load_headers,
    read_csv_rows,
    transform_customers,
    write_workbook,
)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Mews_import_v2.10.0.xlsx"
CUSTOMERS_CSV = ROOT / "contract-extraction" / "V17-customers-export-21-10-2026.csv"
COMPANIES_CSV = ROOT / "contract-extraction" / "V17-companies-export-21-10-2026.csv"
MANIFEST = ROOT / "contract-extraction" / "V17-mews-import-manifest.json"
MIN_ARRIVAL = date(2026, 11, 1)
OUTPUTS = [
    Path.home() / "Downloads" / "V17_Full_Mews_import_V6_B2C_and_B2B_21_10_2026_combined.xlsx",
    ROOT / "Mews_import_V17_Full_V6_B2C_B2B_21_10_2026_combined.xlsx",
]

NOTE_COLS = {"Reservation Notes", "Internal Guest Notes"}
COUNTRY_COLS = ("Nationality Code", "Country Code", "Document Issuing Country Code")
COMPANY_REF_COLS = ("Partner Company", "Company Id")
SEX = {
    "male": "Male",
    "hombre": "Male",
    "masculino": "Male",
    "d": "Male",
    "don": "Male",
    "female": "Female",
    "mujer": "Female",
    "femenino": "Female",
    "dña": "Female",
    "doña": "Female",
}
PLACEHOLDERS = {"address_line_1", "address_line_2"}
IPV4 = re.compile(r"\b\d{1,3}(\.\d{1,3}){3}\b")
MIN_GUEST_AGE_YEARS = 14


def parse_iso(v) -> date | None:
    if not v:
        return None
    try:
        return datetime.strptime(str(v).strip()[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def unwrap(v):
    if not isinstance(v, str):
        return v
    s = v.strip().strip("/").strip()
    return s or None


def clean_company_name(v):
    if v is None:
        return None
    s = re.sub(r"^[\s./]+", "", str(v)).strip()
    return s or None


def append_note(row: list, idx: dict[str, int], text: str) -> None:
    i = idx["Reservation Notes"]
    row[i] = f"{row[i]} | {text}" if row[i] else text


def clean_customer(row: list, idx: dict[str, int], today: date, fixes: dict[str, list]) -> None:
    gc = row[idx["Group Code"]]

    def fix(kind: str, col: str, new) -> None:
        old = row[idx[col]]
        if old != new:
            fixes[kind].append({"group_code": gc, "column": col, "from": old, "to": new})
            row[idx[col]] = new

    for col, i in idx.items():
        if col in NOTE_COLS or not isinstance(row[i], str):
            continue
        new = unwrap(row[i])
        if new and new.lower() in PLACEHOLDERS:
            new = None
        fix("text_cleanup", col, new)

    for col in COUNTRY_COLS:
        v = row[idx[col]]
        if v:
            code = re.sub(r"[^A-Za-z]", "", v).upper()
            fix("country_code", col, code if len(code) == 2 else None)

    if row[idx["Sex"]]:
        fix("sex", "Sex", SEX.get(re.sub(r"[^a-zñ]", "", row[idx["Sex"]].lower())))

    bd = row[idx["Birth Date"]]
    if bd:
        d = parse_iso(bd)
        valid = d is not None and 1900 <= d.year <= today.year - MIN_GUEST_AGE_YEARS
        fix("birth_date", "Birth Date", d.isoformat() if valid else None)

    tel = row[idx["Telephone"]]
    if tel and ("@" in tel or IPV4.search(tel)):
        fix("telephone", "Telephone", None)

    tin = row[idx["Tax Identification Number"]]
    if tin and not any(ch.isdigit() for ch in tin):
        fix("tax_id", "Tax Identification Number", None)

    for col in COMPANY_REF_COLS:
        fix("company_name", col, clean_company_name(row[idx[col]]))


def split_shared_emails(rows: list[list], idx: dict[str, int]) -> dict[str, int]:
    groups: dict[str, list[list]] = defaultdict(list)
    for r in rows:
        e = (r[idx["Email"]] or "").strip().lower()
        if e:
            groups[e].append(r)
    split: dict[str, int] = {}
    for email, rs in groups.items():
        firsts = {(r[idx["First Name"]] or "").strip().lower() for r in rs} - {""}
        docs = {(r[idx["Document Number"]] or "").strip().upper() for r in rs} - {""}
        if len(firsts) > 1 and len(docs) > 1:
            for r in rs:
                append_note(r, idx, f"Booker email: {r[idx['Email']]}")
                r[idx["Email"]] = None
            split[email] = len(rs)
    return split


def clean_companies(rows: list[list], headers: list[str]) -> list[tuple]:
    cidx = {h: i for i, h in enumerate(headers)}
    out: list[tuple] = []
    seen: set[str] = set()
    for r in rows:
        r = [unwrap(v) for v in r]
        name = clean_company_name(r[cidx["Company Name(*)"]])
        if not name or name.lower() in seen:
            continue
        r[cidx["Company Name(*)"]] = name
        seen.add(name.lower())
        out.append(tuple(r))
    return out


def verify(path: Path, min_arrival: date) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Customers"]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    idx = {k: i for i, k in enumerate(h)}
    rows = [r for r in it if any(v not in (None, "") for v in r)]
    co = wb["Companies"]
    cit = co.iter_rows(values_only=True)
    ch = list(next(cit))
    names = {str(r[ch.index("Company Name(*)")]).strip().lower() for r in cit if r[0]}
    wb.close()

    problems: list[str] = []
    rooms: dict[str, list[tuple[date, date, str]]] = defaultdict(list)
    for r in rows:
        gc = r[idx["Group Code"]]
        a, d = parse_iso(r[idx["Arrival Date(*)"]]), parse_iso(r[idx["Departure Date(*)"]])
        if not a or a < min_arrival:
            problems.append(f"{gc}: arrival {r[idx['Arrival Date(*)']]} before {min_arrival.isoformat()}")
        if not d or (a and d <= a):
            problems.append(f"{gc}: departure not after arrival")
        for col in COUNTRY_COLS:
            v = r[idx[col]]
            if v and not re.fullmatch(r"[A-Z]{2}", str(v)):
                problems.append(f"{gc}: {col} {v!r}")
        if r[idx["Sex"]] not in (None, "Male", "Female"):
            problems.append(f"{gc}: Sex {r[idx['Sex']]!r}")
        for col in COMPANY_REF_COLS:
            v = r[idx[col]]
            if v and str(v).strip().lower() not in names:
                problems.append(f"{gc}: {col} {v!r} not on Companies sheet")
        if bool(r[idx["Prepaid Amount"]]) != bool(r[idx["Prepayment Type"]]):
            problems.append(f"{gc}: prepaid amount/type mismatch")
        room = r[idx["Allocated Room Number"]]
        if room and a and d:
            rooms[str(room)].append((a, d, gc))
    for room, stays in rooms.items():
        stays.sort()
        for (a1, d1, g1), (a2, _, g2) in zip(stays, stays[1:]):
            if a2 < d1:
                problems.append(f"room {room}: {g1} overlaps {g2}")
    return problems


def puerta_prepayment_summary(customers_t: list[tuple], idx: dict[str, int]) -> dict:
    puerta = []
    for r in customers_t:
        partner = r[idx["Partner Company"]]
        if partner and "puerta de sevilla" in str(partner).lower():
            puerta.append(
                {
                    "group_code": r[idx["Group Code"]],
                    "prepaid": r[idx["Prepaid Amount"]],
                    "type": r[idx["Prepayment Type"]],
                }
            )
    return {"count": len(puerta), "rows": puerta}


def main() -> None:
    if not CUSTOMERS_CSV.is_file():
        raise SystemExit(f"Missing {CUSTOMERS_CSV}. Run: node contract-extraction/export-v6-tables-csv-21-10-2026.mjs V17")

    tpl = openpyxl.load_workbook(TEMPLATE, read_only=True)
    cust_headers = load_headers(tpl["Customers"])
    co_headers = load_headers(tpl["Companies"])
    tpl.close()
    idx = {h: i for i, h in enumerate(cust_headers)}
    today = date.today()

    customers = [list(r) for r in transform_customers(read_csv_rows(CUSTOMERS_CSV, cust_headers), cust_headers)]
    fixes: dict[str, list] = defaultdict(list)
    for r in customers:
        clean_customer(r, idx, today, fixes)
    shared_emails = split_shared_emails(customers, idx)
    customers_t = [tuple(r) for r in customers]

    companies = clean_companies(read_csv_rows(COMPANIES_CSV, co_headers), co_headers)
    co_subset = companies_for_customers(customers_t, co_headers, companies, cust_headers)

    for out in OUTPUTS:
        write_workbook(out, TEMPLATE, cust_headers, customers_t, co_headers, co_subset)

    problems = verify(OUTPUTS[0], MIN_ARRIVAL)
    puerta = puerta_prepayment_summary(customers_t, idx)
    manifest = {
        "source": "V6_Final_B2C_and_B2B + V6_Final_Companies (Supabase; NC_1147 Adult Count=2)",
        "min_arrival": MIN_ARRIVAL.isoformat(),
        "customers": len(customers_t),
        "companies": len(co_subset),
        "viajes_puerta_de_sevilla_prepayment": puerta,
        "shared_emails_split": shared_emails,
        "field_fixes": {k: v for k, v in fixes.items() if k != "text_cleanup"},
        "text_cleanup_count": len(fixes.get("text_cleanup", [])),
        "verify_problems": problems,
        "outputs": [str(p) for p in OUTPUTS],
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    print(f"V17: {len(customers_t)} customers, {len(co_subset)} companies (V6, arrival >= {MIN_ARRIVAL})")
    print(f"Viajes Puerta de Sevilla rows: {puerta['count']}")
    print(f"verify problems: {len(problems)}")
    for p in problems:
        print(f"  {p}")
    for out in OUTPUTS:
        print(f"-> {out}")


if __name__ == "__main__":
    main()
