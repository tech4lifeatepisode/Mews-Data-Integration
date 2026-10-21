"""V8 random-10: same customers as V7 + Companies sheet (linked V2 companies when present)."""
import random
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Mews_import_v2.10.0.xlsx"
SOURCE = ROOT / "Mews_import_V3_B2C_20_10_2026_combined.xlsx"
COMPANIES_SOURCE = ROOT / "Mews_import_V2_19_10_2026_combined.xlsx"
OUTPUT = Path.home() / "Downloads" / "V8_Random_10_Sample_Mews_import_V3_B2C_20_10_2026_combined.xlsx"

SAMPLE_SIZE = 10
RANDOM_SEED = 20261020
MIN_ARRIVAL_DATE = datetime(2026, 10, 21)
PREPAYMENT_TYPE = "Deposit"
MONTHLY_AUDIT_RE = re.compile(r"Mews import audit — Monthly rent EUR:", re.I)


def header_row(ws) -> list[str]:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def read_data_rows(ws) -> list[tuple]:
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            rows.append(row)
    return rows


def parse_date(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return None


def stay_days(row, idx: dict[str, int]) -> int | None:
    a = parse_date(row[idx["Arrival Date(*)"]])
    d = parse_date(row[idx["Departure Date(*)"]])
    if not a or not d:
        return None
    return (d - a).days


def row_tags(row, idx: dict[str, int]) -> set[str]:
    tags = set()
    cat = row[idx["Allocated Room Category(*)"]]
    if cat:
        tags.add(f"cat:{cat}")
    tags.add("products:yes" if row[idx["Included Products"]] else "products:no")
    days = stay_days(row, idx)
    if days is not None:
        if days < 90:
            tags.add("stay:short")
        elif days <= 270:
            tags.add("stay:medium")
        else:
            tags.add("stay:long")
    return tags


def score_row(row, idx: dict[str, int], covered: set[str]) -> int:
    tags = row_tags(row, idx)
    novel = tags - covered
    score = len(novel) * 10
    if row[idx["Included Products"]]:
        score += 5
    days = stay_days(row, idx)
    if days is not None:
        score += min(days // 30, 12)
    return score


def pick_diverse_sample(rows: list[tuple], headers: list[str], n: int) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    rng = random.Random(RANDOM_SEED)
    pool = rows[:]
    rng.shuffle(pool)

    selected: list[tuple] = []
    covered: set[str] = set()
    used: set[int] = set()

    by_cat: dict[str, list[tuple]] = defaultdict(list)
    for row in pool:
        cat = row[idx["Allocated Room Category(*)"]] or "(blank)"
        by_cat[str(cat)].append(row)

    for cat in sorted(by_cat.keys(), key=lambda c: (-len(by_cat[c]), c)):
        if len(selected) >= n:
            break
        for row in sorted(by_cat[cat], key=lambda r: -score_row(r, idx, covered)):
            if id(row) in used:
                continue
            selected.append(row)
            used.add(id(row))
            covered |= row_tags(row, idx)
            break

    quotas = [("products:yes", 4), ("products:no", 4), ("stay:short", 2), ("stay:medium", 4), ("stay:long", 2)]
    counts = {k: sum(1 for r in selected if k in row_tags(r, idx)) for k, _ in quotas}

    while len(selected) < n:
        best, best_sc = None, -1.0
        for row in pool:
            if id(row) in used:
                continue
            tags = row_tags(row, idx)
            bonus = sum(15 for key, target in quotas if key in tags and counts[key] < target)
            sc = bonus + score_row(row, idx, covered) + rng.random()
            if sc > best_sc:
                best, best_sc = row, sc
        if best is None:
            break
        selected.append(best)
        used.add(id(best))
        covered |= row_tags(best, idx)
        for key, _ in quotas:
            if key in row_tags(best, idx):
                counts[key] += 1

    return selected[:n]


def parse_money(val) -> str | None:
    if val is None or val == "":
        return None
    s = str(val).strip()
    if "Stays" in s:
        return None
    m = re.match(r"^(\d+(?:\.\d+)?)", s.replace(",", ".").replace(" ", ""))
    if not m:
        return None
    n = float(m.group(1))
    return str(int(n)) if n == int(n) else str(n)


def monthly_from_row(row: list, idx: dict[str, int]) -> str | None:
    notes_i = idx.get("Reservation Notes")
    if notes_i is not None and row[notes_i]:
        notes_s = str(row[notes_i])
        for pat in (
            r"Mews import audit — Monthly rent EUR:\s*([\d.,]+)",
            r"Monthly price:\s*([\d.,]+)",
            r"Monthly rent EUR:\s*([\d.,]+)",
        ):
            m = re.search(pat, notes_s, re.I)
            if m:
                return parse_money(m.group(1))
    if "Per Night Prices" in idx:
        m = parse_money(row[idx["Per Night Prices"]])
        if m:
            return m
    return None


def append_audit_notes(row: list, idx: dict[str, int]) -> None:
    notes_i = idx.get("Reservation Notes")
    rate_i = idx.get("Allocated Rate")
    if notes_i is None:
        return
    notes = row[notes_i]
    notes_s = "" if notes is None else str(notes).strip()
    if MONTHLY_AUDIT_RE.search(notes_s):
        return
    monthly = monthly_from_row(row, idx)
    if not monthly:
        return
    rate = row[rate_i] if rate_i is not None else ""
    audit = (
        f"Mews import audit — Monthly rent EUR: {monthly}; "
        f"Per Night Prices left empty; Exact rate: {rate or 'n/a'}"
    )
    row[notes_i] = f"{notes_s} | {audit}" if notes_s else audit


def stash_companions_in_notes(row: list, idx: dict[str, int]) -> None:
    comp_i = idx.get("Companions")
    notes_i = idx.get("Reservation Notes")
    if comp_i is None or notes_i is None:
        return
    companions = row[comp_i]
    if companions is None or str(companions).strip() == "":
        return
    name = str(companions).strip()
    notes = row[notes_i]
    notes_s = "" if notes is None else str(notes).strip()
    tag = f"Companion(s) (not imported - add in Mews or re-import with customer ID): {name}"
    if name not in notes_s:
        row[notes_i] = f"{notes_s} | {tag}" if notes_s else tag
    row[comp_i] = None


def mews_import_ready(row: list, idx: dict[str, int]) -> None:
    stash_companions_in_notes(row, idx)
    append_audit_notes(row, idx)
    if "Person Counts(*)" in idx:
        row[idx["Person Counts(*)"]] = None
    if "Per Night Prices" in idx:
        row[idx["Per Night Prices"]] = None
    if "Total Price" in idx:
        row[idx["Total Price"]] = None
    if "Arrival Time" in idx:
        row[idx["Arrival Time"]] = None
    if "Departure Time" in idx:
        row[idx["Departure Time"]] = None
    prepaid_i = idx.get("Prepaid Amount")
    prep_type_i = idx.get("Prepayment Type")
    if prepaid_i is None or prep_type_i is None:
        return
    prepaid = row[prepaid_i]
    if prepaid is not None and str(prepaid).strip() and any(ch.isdigit() for ch in str(prepaid)):
        row[prep_type_i] = PREPAYMENT_TYPE
    else:
        row[prepaid_i] = None
        row[prep_type_i] = None


def transform_row(row: tuple, idx: dict[str, int]) -> tuple:
    row = list(row)
    mews_import_ready(row, idx)
    return tuple(row)


def company_matches_customer(company_row, cust_row, cidx: dict[str, int], uidx: dict[str, int]) -> bool:
    name = company_row[cidx["Company Name(*)"]]
    if not name:
        return False
    name_s = str(name).strip().lower()
    for key in ("Partner Company", "Company Id"):
        if key not in uidx:
            continue
        v = cust_row[uidx[key]]
        if v and str(v).strip().lower() == name_s:
            return True
    tax_i = cidx.get("Tax ID Number")
    if tax_i is not None and company_row[tax_i]:
        t = str(company_row[tax_i]).strip().lower()
        for key in ("Tax Identification Number", "Document Number"):
            if key in uidx:
                v = cust_row[uidx[key]]
                if v and t in str(v).lower():
                    return True
    return False


def companies_for_customers(
    picked: list[tuple], co_headers: list[str], all_companies: list[tuple], cust_headers: list[str]
) -> list[tuple]:
    if not all_companies:
        return []
    cidx = {h: i for i, h in enumerate(co_headers)}
    uidx = {h: i for i, h in enumerate(cust_headers)}
    picked_companies: list[tuple] = []
    seen: set[str] = set()
    for crow in all_companies:
        key = crow[cidx["Company Name(*)"]]
        if not key or key in seen:
            continue
        if any(company_matches_customer(crow, u, cidx, uidx) for u in picked):
            picked_companies.append(crow)
            seen.add(key)
    return picked_companies


def write_sheet(ws, headers: list[str], data_rows: list[tuple]) -> None:
    for c, h in enumerate(headers, start=1):
        if ws.cell(row=1, column=c).value != h:
            raise SystemExit(f"Header mismatch col {c}: {ws.cell(row=1, column=c).value!r} != {h!r}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    room_col = headers.index("Allocated Room Number") + 1 if "Allocated Room Number" in headers else None
    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if room_col == c_idx and val is not None and str(val).strip().startswith("-"):
                cell.number_format = "@"
                cell.value = str(val).strip()


def main() -> None:
    if not SOURCE.is_file():
        raise SystemExit(f"Source not found: {SOURCE}. Run build-v3-mews-import-xlsx.mjs first.")
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    src_wb = openpyxl.load_workbook(SOURCE, read_only=True)
    cust_headers = header_row(src_wb["Customers"])
    co_headers = header_row(src_wb["Companies"])
    all_customers = read_data_rows(src_wb["Customers"])
    src_wb.close()

    all_companies: list[tuple] = []
    if COMPANIES_SOURCE.is_file():
        co_wb = openpyxl.load_workbook(COMPANIES_SOURCE, read_only=True)
        co_headers = header_row(co_wb["Companies"])
        all_companies = read_data_rows(co_wb["Companies"])
        co_wb.close()

    uidx = {h: i for i, h in enumerate(cust_headers)}

    def arrival_on_or_after_min(row: tuple) -> bool:
        a = parse_date(row[uidx["Arrival Date(*)"]])
        return bool(a and a >= MIN_ARRIVAL_DATE)

    future_customers = [r for r in all_customers if arrival_on_or_after_min(r)]
    pool = future_customers if len(future_customers) >= SAMPLE_SIZE else all_customers
    if pool is future_customers:
        print(f"  Using {len(future_customers)} rows with arrival >= {MIN_ARRIVAL_DATE.date()}")

    picked = [transform_row(r, uidx) for r in pick_diverse_sample(pool, cust_headers, SAMPLE_SIZE)]

    shutil.copy2(TEMPLATE, OUTPUT)
    out_wb = openpyxl.load_workbook(OUTPUT)
    write_sheet(out_wb["Customers"], cust_headers, picked)
    picked_companies = companies_for_customers(picked, co_headers, all_companies, cust_headers)
    write_sheet(out_wb["Companies"], co_headers, picked_companies)
    out_wb.save(OUTPUT)

    prep = sum(
        1
        for r in picked
        if uidx.get("Prepayment Type") is not None and r[uidx["Prepayment Type"]] == PREPAYMENT_TYPE
    )
    person_empty = sum(1 for r in picked if not r[uidx["Person Counts(*)"]])
    print(f"Wrote {OUTPUT}")
    print(f"  Customers: {len(picked)}")
    print(f"  Companies: {len(picked_companies)}")
    print(f"  Prepayment Type Deposit: {prep}/{len(picked)}")
    print(f"  Person Counts(*) empty: {person_empty}/{len(picked)}")
    for r in picked:
        email = r[uidx["Email"]]
        room = r[uidx["Allocated Room Number"]]
        cat = r[uidx["Allocated Room Category(*)"]]
        print(f"    {email} | {cat} | room {room}")