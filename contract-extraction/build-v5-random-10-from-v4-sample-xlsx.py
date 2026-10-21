"""V5 random-10: V4 Mews fixes + audit monthly rent in Reservation Notes + Prepayment Type = Deposit."""
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Mews_import_v2.10.0.xlsx"
SOURCE = Path.home() / "Downloads" / "V4_Random_10_Sample_Mews_import_V3_B2C_20_10_2026_combined.xlsx"
OUTPUT = Path.home() / "Downloads" / "V5_Random_10_Sample_Mews_import_V3_B2C_20_10_2026_combined.xlsx"

PREPAYMENT_TYPE = "Deposit"
MONTHLY_AUDIT_RE = re.compile(r"Mews import audit — Monthly rent EUR:", re.I)


def parse_money(val) -> str | None:
    if val is None or val == "":
        return None
    s = str(val).strip()
    if "Stays" in s:
        return None
    m = re.match(r"^(\d+(?:\.\d+)?)", s.replace(",", ".").replace(" ", ""))
    if not m:
        return None
    n = float(m.group(1))
    return str(int(n)) if n == int(n) else str(n)


def monthly_from_row(row: list, idx: dict[str, int]) -> str | None:
    notes_i = idx.get("Reservation Notes")
    if notes_i is not None and row[notes_i]:
        notes_s = str(row[notes_i])
        for pat in (
            r"Mews import audit — Monthly rent EUR:\s*([\d.,]+)",
            r"Monthly price:\s*([\d.,]+)",
            r"Monthly rent EUR:\s*([\d.,]+)",
        ):
            m = re.search(pat, notes_s, re.I)
            if m:
                return parse_money(m.group(1))
    if "Per Night Prices" in idx:
        m = parse_money(row[idx["Per Night Prices"]])
        if m:
            return m
    return None


def append_audit_notes(row: list, idx: dict[str, int]) -> None:
    notes_i = idx.get("Reservation Notes")
    rate_i = idx.get("Allocated Rate")
    if notes_i is None:
        return
    notes = row[notes_i]
    notes_s = "" if notes is None else str(notes).strip()
    if MONTHLY_AUDIT_RE.search(notes_s):
        return
    monthly = monthly_from_row(row, idx)
    if not monthly:
        return
    rate = row[rate_i] if rate_i is not None else ""
    audit = (
        f"Mews import audit — Monthly rent EUR: {monthly}; "
        f"Per Night Prices left empty; Exact rate: {rate or 'n/a'}"
    )
    row[notes_i] = f"{notes_s} | {audit}" if notes_s else audit


def mews_import_ready(row: list, idx: dict[str, int]) -> None:
    append_audit_notes(row, idx)
    if "Person Counts(*)" in idx:
        row[idx["Person Counts(*)"]] = None
    if "Per Night Prices" in idx:
        row[idx["Per Night Prices"]] = None
    if "Total Price" in idx:
        row[idx["Total Price"]] = None
    if "Arrival Time" in idx:
        row[idx["Arrival Time"]] = None
    if "Departure Time" in idx:
        row[idx["Departure Time"]] = None
    prepaid_i = idx.get("Prepaid Amount")
    prep_type_i = idx.get("Prepayment Type")
    if prepaid_i is None or prep_type_i is None:
        return
    prepaid = row[prepaid_i]
    if prepaid is not None and str(prepaid).strip() and any(ch.isdigit() for ch in str(prepaid)):
        row[prep_type_i] = PREPAYMENT_TYPE
    else:
        row[prepaid_i] = None
        row[prep_type_i] = None


def transform_row(row: tuple, idx: dict[str, int]) -> tuple:
    row = list(row)
    mews_import_ready(row, idx)
    return tuple(row)


def main() -> None:
    if not SOURCE.is_file():
        raise SystemExit(f"Missing source: {SOURCE}")
    src = openpyxl.load_workbook(SOURCE, read_only=True)
    headers = [c.value for c in next(src["Customers"].iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in src["Customers"].iter_rows(min_row=2, values_only=True):
        if any(cell is not None and str(cell).strip() for cell in row):
            rows.append(transform_row(row, idx))
    src.close()

    shutil.copy2(TEMPLATE, OUTPUT)
    wb = openpyxl.load_workbook(OUTPUT)
    ws = wb["Customers"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    ws_co = wb["Companies"]
    if ws_co.max_row > 1:
        ws_co.delete_rows(2, ws_co.max_row - 1)
    wb.save(OUTPUT)
    prep = sum(1 for r in rows if idx.get("Prepayment Type") is not None and r[idx["Prepayment Type"]] == PREPAYMENT_TYPE)
    print(f"Wrote {OUTPUT} ({len(rows)} customers, Prepayment Type Deposit: {prep})")


if __name__ == "__main__":
    main()
