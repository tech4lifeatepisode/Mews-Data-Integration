"""Copy Mews template xlsx and fill Customers + Companies sheets from CSV exports."""
import csv
import shutil
import sys
from pathlib import Path

import openpyxl

CUSTOMER_HEADERS = [
    "Email",
    "Last Name(*)",
    "First Name",
    "Second Last Name",
    "Title",
    "Sex",
    "Nationality Code",
    "Preferred Language Code",
    "Birth Date",
    "Birth Place",
    "Telephone",
    "Document Type",
    "Document Number",
    "Document Expiry Date",
    "Document Issue Date",
    "Document Issuing Country Code",
    "Document Issuing Country Subdivision Code",
    "Document Issuing City",
    "Document Support Number",
    "Occupation",
    "Car Registration Number",
    "Address Line 1",
    "Address Line 2",
    "City",
    "Postal Code",
    "Country Code",
    "Country Subdivision Code",
    "Customer Classification",
    "Internal Guest Notes",
    "Preferred Space Features",
    "Dietary Requirements",
    "Loyalty Code",
    "Company Id",
    "Tax Identification Number",
    "Accounting Code",
    "Billing Code",
    "Italian Destination Code",
    "Italian Fiscal Code",
    "Italian Lottery Code",
    "Invoiceable",
    "Process ID Image",
    "Process Verification Photo",
    "Send Marketing Emails",
    "Send Marketing Postal Mail",
    "Send Partner Marketing Emails",
    "Send Partner Marketing Postal Mail",
    "Arrival Date(*)",
    "Arrival Time",
    "Departure Date(*)",
    "Departure Time",
    "Person Counts(*)",
    "Adult Count(*)",
    "Child Count(*)",
    "Booker ID",
    "Companions",
    "Reservation Notes",
    "Confirmation Code",
    "Group Code",
    "Voucher Code",
    "Availability Block ID",
    "Currency Code",
    "Total Price",
    "Per Night Prices",
    "Prepaid Amount",
    "Prepayment Type",
    "Skip Automatic Settlement",
    "Allocated Room Category(*)",
    "Allocated Room Number",
    "Allocated Rate",
    "Included Products",
    "Partner Company",
    "Travel Agency",
    "Business Segment",
    "Reservation Purpose",
    "Release Date",
    "Channel manager number",
    "Channel number",
    "Channel manager integration ID",
    "Source",
]

COMPANY_HEADERS = [
    "Company Name(*)",
    "Telephone",
    "Contact Person",
    "Contact Notes",
    "Website",
    "Notes",
    "Department",
    "Address Line 1",
    "Address Line 2",
    "City",
    "Postal Code",
    "Country Code",
    "Mother Company Id",
    "Reference ID",
    "Fiscal Identifier",
    "Tax ID Number",
    "Additional Tax Identifier",
    "Accounting Code",
    "Billing Code",
    "External Identifier",
    "Invoice Due Interval",
    "Invoicing Email",
    "Invoiceable",
    "Add Fees To Invoices",
    "Add Tax Deducted Payment To Invoices",
    "IATA",
    "DUNS Number",
    "Cod. destinatario",
    "Fiscal code",
    "Tender Code (CIG)",
    "Project Code (CUP)",
]

SKIP_COLUMNS = {"id", "created_at", "nc_number"}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def cell_value(raw: str | None):
    if raw is None:
        return None
    s = str(raw).strip()
    if s == "":
        return None
    return s


def assert_template_headers(ws, expected: list[str]) -> None:
    actual = [ws.cell(row=1, column=c).value for c in range(1, len(expected) + 1)]
    if actual != expected:
        raise SystemExit(
            f"Template header mismatch on sheet {ws.title!r}.\n"
            f"Expected {len(expected)} columns matching Mews_import_v2.10.0.xlsx.\n"
            f"First diff at col {next((i for i, (a, e) in enumerate(zip(actual, expected)) if a != e), '?')}: "
            f"{actual!r} vs {expected!r}"
        )


def write_sheet(ws, headers: list[str], rows: list[dict[str, str]]) -> int:
    assert_template_headers(ws, headers)
    # Clear prior data rows (template reserves empty rows)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, key in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c_idx, value=cell_value(row.get(key)))
    return len(rows)


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit(
            "Usage: python _build_v2_mews_import_xlsx.py <template.xlsx> <customers.csv> <companies.csv> <out.xlsx>"
        )
    template, customers_csv, companies_csv, out_xlsx = map(Path, sys.argv[1:5])
    if not template.is_file():
        raise SystemExit(f"Template not found: {template}")
    if not customers_csv.is_file():
        raise SystemExit(f"Customers CSV not found: {customers_csv}")
    if not companies_csv.is_file():
        raise SystemExit(f"Companies CSV not found: {companies_csv}")

    customer_rows = read_csv(customers_csv)
    company_rows = read_csv(companies_csv)

    for label, rows, headers in (
        ("customers", customer_rows, CUSTOMER_HEADERS),
        ("companies", company_rows, COMPANY_HEADERS),
    ):
        extra = set(rows[0].keys()) - set(headers) - SKIP_COLUMNS if rows else set()
        missing = set(headers) - set(rows[0].keys()) - SKIP_COLUMNS if rows else set()
        if missing:
            raise SystemExit(f"{label} CSV missing columns: {sorted(missing)}")
        if extra:
            print(f"Note: ignoring extra {label} CSV columns: {sorted(extra)}")

    shutil.copy2(template, out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    n_cust = write_sheet(wb["Customers"], CUSTOMER_HEADERS, customer_rows)
    n_co = write_sheet(wb["Companies"], COMPANY_HEADERS, company_rows)
    wb.save(out_xlsx)
    print(f"Wrote {out_xlsx}")
    print(f"  Customers: {n_cust} rows")
    print(f"  Companies: {n_co} rows")


if __name__ == "__main__":
    main()
