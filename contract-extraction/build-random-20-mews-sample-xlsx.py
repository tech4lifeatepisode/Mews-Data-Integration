"""Build a diverse 20-row Customers sample + matching Companies from V2 combined xlsx."""
import random
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Mews_import_v2.10.0.xlsx"
SOURCE = ROOT / "Mews_import_V2_19_10_2026_combined.xlsx"
OUTPUT = ROOT / "Random_20_Sample_Mews_import_V2_19_10_2026_combined.xlsx"

SAMPLE_SIZE = 20
RANDOM_SEED = 20261020


def header_row(ws) -> list[str]:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def read_data_rows(ws) -> list[tuple]:
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            rows.append(row)
    return rows


def parse_date(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return None


def stay_days(row, idx: dict[str, int]) -> int | None:
    a = parse_date(row[idx["Arrival Date(*)"]])
    d = parse_date(row[idx["Departure Date(*)"]])
    if not a or not d:
        return None
    return (d - a).days


def stay_bucket(days: int | None) -> str:
    if days is None:
        return "unknown"
    if days < 90:
        return "short"
    if days <= 270:
        return "medium"
    return "long"


def row_tags(row, idx: dict[str, int]) -> set[str]:
    tags = set()
    cat = row[idx["Allocated Room Category(*)"]]
    if cat:
        tags.add(f"cat:{cat}")
    tags.add("products:yes" if row[idx["Included Products"]] else "products:no")
    tags.add("b2b:yes" if row[idx["Partner Company"]] else "b2b:no")
    tags.add(f"stay:{stay_bucket(stay_days(row, idx))}")
    seg = row[idx.get("Business Segment", -1)] if "Business Segment" in idx else None
    if seg and str(seg).strip().upper() == "B2B":
        tags.add("segment:b2b")
    if row[idx["Company Id"]] and not row[idx["Partner Company"]]:
        tags.add("company_id_only")
    return tags


def score_row(row, idx: dict[str, int], covered: set[str]) -> int:
    tags = row_tags(row, idx)
    novel = tags - covered
    score = len(novel) * 10
    if row[idx["Included Products"]]:
        score += 2
    if row[idx["Partner Company"]]:
        score += 3
    days = stay_days(row, idx)
    if days is not None:
        score += min(days // 30, 12)
    return score


def pick_diverse_sample(rows: list[tuple], headers: list[str], n: int) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    rng = random.Random(RANDOM_SEED)
    pool = rows[:]
    rng.shuffle(pool)

    selected: list[tuple] = []
    covered: set[str] = set()
    used_indices: set[int] = set()

    # Phase 1: ensure at least one row per major room category present in data
    by_cat: dict[str, list[tuple]] = defaultdict(list)
    for row in pool:
        cat = row[idx["Allocated Room Category(*)"]] or "(blank)"
        by_cat[str(cat)].append(row)

    priority_cats = sorted(by_cat.keys(), key=lambda c: (-len(by_cat[c]), c))
    for cat in priority_cats:
        if len(selected) >= n:
            break
        if len(by_cat[cat]) > len(priority_cats) and cat in (
            c for c in priority_cats[:8]
        ):
            pass
        candidates = sorted(by_cat[cat], key=lambda r: -score_row(r, idx, covered))
        for row in candidates:
            rid = id(row)
            if rid in used_indices:
                continue
            selected.append(row)
            used_indices.add(rid)
            covered |= row_tags(row, idx)
            break

    # Phase 2: fill quotas for products / b2b / stay length
    quotas = [
        ("products:yes", 5),
        ("products:no", 5),
        ("b2b:yes", 4),
        ("b2b:no", 8),
        ("stay:short", 4),
        ("stay:medium", 8),
        ("stay:long", 6),
    ]
    counts = {k: sum(1 for r in selected if k in row_tags(r, idx)) for k, _ in quotas}

    while len(selected) < n:
        best = None
        best_score = -1
        for row in pool:
            if id(row) in used_indices:
                continue
            tags = row_tags(row, idx)
            q_bonus = 0
            for key, target in quotas:
                if key in tags and counts[key] < target:
                    q_bonus += 15
            sc = q_bonus + score_row(row, idx, covered) + rng.random()
            if sc > best_score:
                best_score = sc
                best = row
        if best is None:
            break
        selected.append(best)
        used_indices.add(id(best))
        covered |= row_tags(best, idx)
        for key, _ in quotas:
            if key in row_tags(best, idx):
                counts[key] += 1

    return selected[:n]


def company_matches_customer(company_row, cust_row, cidx: dict, uidx: dict) -> bool:
    name = company_row[cidx["Company Name(*)"]]
    if not name:
        return False
    name_s = str(name).strip().lower()
    for key in ("Partner Company", "Company Id"):
        v = cust_row[uidx[key]]
        if v and str(v).strip().lower() == name_s:
            return True
    tax = company_row[cidx.get("Tax ID Number", -1)]
    if tax:
        t = str(tax).strip().lower()
        for key in ("Tax Identification Number", "Document Number"):
            if key in uidx:
                v = cust_row[uidx[key]]
                if v and t in str(v).lower():
                    return True
    return False


def write_sheet(ws, headers: list[str], data_rows: list[tuple]) -> None:
    for c, h in enumerate(headers, start=1):
        if ws.cell(row=1, column=c).value != h:
            raise SystemExit(f"Header mismatch col {c}: {ws.cell(row=1, column=c).value!r} != {h!r}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)


def main() -> None:
    if not SOURCE.is_file():
        raise SystemExit(f"Source not found: {SOURCE}")
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    src_wb = openpyxl.load_workbook(SOURCE, read_only=True)
    cust_headers = header_row(src_wb["Customers"])
    co_headers = header_row(src_wb["Companies"])
    all_customers = read_data_rows(src_wb["Customers"])
    all_companies = read_data_rows(src_wb["Companies"])
    src_wb.close()

    picked = pick_diverse_sample(all_customers, cust_headers, SAMPLE_SIZE)
    cidx = {h: i for i, h in enumerate(co_headers)}
    uidx = {h: i for i, h in enumerate(cust_headers)}

    picked_companies = []
    seen_co = set()
    for crow in all_companies:
        key = crow[cidx["Company Name(*)"]]
        if not key or key in seen_co:
            continue
        if any(company_matches_customer(crow, u, cidx, uidx) for u in picked):
            picked_companies.append(crow)
            seen_co.add(key)

    shutil.copy2(TEMPLATE, OUTPUT)
    out_wb = openpyxl.load_workbook(OUTPUT)
    write_sheet(out_wb["Customers"], cust_headers, picked)
    write_sheet(out_wb["Companies"], co_headers, picked_companies)
    out_wb.save(OUTPUT)

    uidx = {h: i for i, h in enumerate(cust_headers)}
    print(f"Wrote {OUTPUT}")
    print(f"  Customers: {len(picked)}")
    print(f"  Companies: {len(picked_companies)}")
    print("  Sample breakdown:")
    cats = defaultdict(int)
    for r in picked:
        cats[r[uidx["Allocated Room Category(*)"]]] += 1
    for cat, n in sorted(cats.items(), key=lambda x: -x[1]):
        print(f"    {cat}: {n}")
    print(
        f"    with Included Products: {sum(1 for r in picked if r[uidx['Included Products']])}"
    )
    print(f"    with Partner Company: {sum(1 for r in picked if r[uidx['Partner Company']])}")
    stays = [stay_days(r, uidx) for r in picked]
    print(f"    stay days min/median/max: {min(stays)}/{sorted(stays)[len(stays)//2]}/{max(stays)}")


if __name__ == "__main__":
    main()
