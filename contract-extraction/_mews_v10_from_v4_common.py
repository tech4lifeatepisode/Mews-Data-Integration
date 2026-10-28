"""Shared V10 Mews xlsx build: sanitize customers, filter companies, write sheets."""
from __future__ import annotations

import csv
import importlib.util
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Callable

import openpyxl

_v8_path = Path(__file__).resolve().parent / "build-v8-random-10-mews-import-sample.py"
_spec = importlib.util.spec_from_file_location("v8_sample", _v8_path)
_v8 = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_v8)

PREPAYMENT_TYPE = _v8.PREPAYMENT_TYPE
mews_import_ready = _v8.mews_import_ready
company_matches_customer = _v8.company_matches_customer

CATEGORY_ALIASES = {
    "2 bedroom (double / single)": "The Two Bedroom (Double Single)",
    "2 bedroom (double/single)": "The Two Bedroom (Double Single)",
    "2 bedroom (2+1)": "The Two Bedroom (Double Single)",
    "the two bedroom (double single)": "The Two Bedroom (Double Single)",
}

INVALID_DOC_TYPES = {"taxidentificationnumber", "tax identification number"}


def parse_date(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return None


def canonicalize_category(raw) -> str | None:
    if raw is None or str(raw).strip() == "":
        return None
    s = str(raw).strip()
    key = s.lower()
    if key in CATEGORY_ALIASES:
        return CATEGORY_ALIASES[key]
    for alias, canon in CATEGORY_ALIASES.items():
        if alias in key:
            return canon
    return s


def sanitize_mews_row(row: list, idx: dict[str, int]) -> None:
    mews_import_ready(row, idx)
    for key in (
        "Source",
        "Business Segment",
        "Channel manager integration ID",
        "Channel manager number",
        "Channel number",
    ):
        if key in idx:
            row[idx[key]] = None

    if "Document Type" in idx:
        dt = row[idx["Document Type"]]
        if dt and str(dt).strip().lower().replace(" ", "") in INVALID_DOC_TYPES:
            doc = row[idx.get("Document Number", -1)] if "Document Number" in idx else ""
            doc_s = str(doc or "").strip()
            row[idx["Document Type"]] = (
                "Passport" if doc_s and doc_s[0].isalpha() and len(doc_s) >= 6 else "IdentityCard"
            )

    if "Included Products" in idx:
        row[idx["Included Products"]] = None

    if "Allocated Room Category(*)" in idx and row[idx["Allocated Room Category(*)"]]:
        canon = canonicalize_category(row[idx["Allocated Room Category(*)"]])
        if canon:
            row[idx["Allocated Room Category(*)"]] = canon


def load_headers(ws) -> list[str]:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def read_csv_rows(path: Path, headers: list[str]) -> list[list]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [[rec.get(h) or None for h in headers] for rec in csv.DictReader(f)]


def transform_customers(rows: list[list], headers: list[str]) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    out: list[tuple] = []
    for row in rows:
        r = list(row)
        sanitize_mews_row(r, idx)
        out.append(tuple(r))
    return out


def companies_for_customers(
    customers: list[tuple], co_headers: list[str], all_companies: list[tuple], cust_headers: list[str]
) -> list[tuple]:
    if not all_companies:
        return []
    cidx = {h: i for i, h in enumerate(co_headers)}
    uidx = {h: i for i, h in enumerate(cust_headers)}
    picked: list[tuple] = []
    seen: set[str] = set()
    for crow in all_companies:
        name = crow[cidx["Company Name(*)"]]
        if not name:
            continue
        key = str(name).strip().lower()
        if key in seen:
            continue
        if any(company_matches_customer(crow, u, cidx, uidx) for u in customers):
            picked.append(crow)
            seen.add(key)
    return picked


def write_sheet(ws, headers: list[str], data_rows: list[tuple]) -> None:
    for c, h in enumerate(headers, start=1):
        if ws.cell(row=1, column=c).value != h:
            raise SystemExit(f"Header mismatch col {c}: {ws.cell(row=1, column=c).value!r} != {h!r}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    room_col = headers.index("Allocated Room Number") + 1 if "Allocated Room Number" in headers else None
    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if room_col == c_idx and val is not None and str(val).strip().startswith("-"):
                cell.number_format = "@"
                cell.value = str(val).strip()


def write_workbook(
    path: Path,
    template: Path,
    cust_headers: list[str],
    customers: list[tuple],
    co_headers: list[str],
    companies: list[tuple],
) -> None:
    shutil.copy2(template, path)
    wb = openpyxl.load_workbook(path)
    write_sheet(wb["Customers"], cust_headers, customers)
    write_sheet(wb["Companies"], co_headers, companies)
    wb.save(path)


def filter_by_arrival_gte(customers: list[tuple], headers: list[str], min_date: datetime) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    arr_i = idx["Arrival Date(*)"]
    return [r for r in customers if (a := parse_date(r[arr_i])) and a >= min_date]


def filter_by_departure_gte(customers: list[tuple], headers: list[str], min_iso: str) -> list[tuple]:
    idx = {h: i for i, h in enumerate(headers)}
    dep_i = idx["Departure Date(*)"]
    kept = []
    for r in customers:
        d = parse_date(r[dep_i])
        if d and d.strftime("%Y-%m-%d") >= min_iso:
            kept.append(r)
    return kept


def build_v10_packages(
    *,
    template: Path,
    customers_csv: Path,
    companies_csv: Path,
    outputs: list[tuple[str, Path, Callable[[list[tuple], list[str]], list[tuple]]]],
    manifest_path: Path | None = None,
) -> dict:
    tpl_wb = openpyxl.load_workbook(template, read_only=True)
    cust_headers = load_headers(tpl_wb["Customers"])
    co_headers = load_headers(tpl_wb["Companies"])
    tpl_wb.close()

    raw = read_csv_rows(customers_csv, cust_headers)
    all_companies = [tuple(r) for r in read_csv_rows(companies_csv, co_headers)]
    customers_all = transform_customers(raw, cust_headers)

    results = {"packages": [], "sanitizedCustomerRows": len(customers_all)}
    for label, out_path, filt in outputs:
        subset = filt(customers_all, cust_headers)
        co_subset = companies_for_customers(subset, co_headers, all_companies, cust_headers)
        write_workbook(out_path, template, cust_headers, subset, co_headers, co_subset)
        idx = {h: i for i, h in enumerate(cust_headers)}
        prep = sum(
            1
            for r in subset
            if idx.get("Prepayment Type") is not None and r[idx["Prepayment Type"]] == PREPAYMENT_TYPE
        )
        pkg = {
            "label": label,
            "path": str(out_path),
            "customers": len(subset),
            "companies": len(co_subset),
            "prepaymentDeposit": prep,
        }
        results["packages"].append(pkg)

    if manifest_path:
        manifest_path.write_text(json.dumps(results, indent=2), encoding="utf-8")
    return results
