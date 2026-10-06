"""
Extract rows from NO SHOW sections in LIBRO workbook files into one sheet.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl
LIBRO_DIR = Path(
    r"c:\Users\Claudio\Desktop\Node Living\Invocies comparison to Guest Report 24.12.25 to 24.04.26\LIBRO"
)

MES_ES = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}


def norm_text(s: object) -> str:
    if s is None:
        return ""
    t = str(s).strip()
    t = unicodedata.normalize("NFD", t)
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return t.upper()


def is_nombre_huesped_header(val: object) -> bool:
    n = norm_text(val)
    return "NOMBRE" in n and "HUESPED" in n


def looks_like_no_show_banner(val: object) -> bool:
    if val is None:
        return False
    t = str(val).strip().upper()
    return "NO SHOW" in t


def row_has_checkout_banner(ws: openpyxl.worksheet.worksheet.Worksheet, r: int, max_c: int) -> bool:
    for c in range(1, max_c + 1):
        n = norm_text(ws.cell(r, c).value)
        if "CHECK" in n and "OUT" in n:
            return True
    return False


def match_header(h: object, *keywords: str) -> bool:
    n = norm_text(h)
    return all(k in n for k in keywords)


def parse_sheet_name_date(name: str) -> datetime | None:
    s = name.strip()
    if s.isdigit():
        if len(s) == 8:
            d, m, y = int(s[:2]), int(s[2:4]), int(s[4:])
            try:
                return datetime(y, m, d)
            except ValueError:
                return None
        if len(s) == 6:
            d, m, y2 = int(s[:2]), int(s[2:4]), int(s[4:])
            try:
                return datetime(2000 + y2, m, d)
            except ValueError:
                return None
    if re.fullmatch(r"\d{1,2}\d{2}", s):
        d, m = int(s[:-2]), int(s[-2:])
        if 1 <= m <= 12:
            return None
    return None


def parse_spanish_date_row(cell_val: object, default_year: int | None) -> datetime | None:
    if cell_val is None:
        return None
    s = str(cell_val).strip()
    m = re.search(
        r"(\d{1,2})\s+de\s+([a-záéíóúñ]+)",
        s,
        flags=re.IGNORECASE,
    )
    if not m:
        return None
    day = int(m.group(1))
    month_name = m.group(2).lower()
    month_name = "".join(
        c for c in unicodedata.normalize("NFD", month_name) if unicodedata.category(c) != "Mn"
    )
    mo = MES_ES.get(month_name)
    if not mo:
        return None
    year = default_year
    if year is None:
        return None
    try:
        return datetime(year, mo, day)
    except ValueError:
        return None


def year_from_workbook_filename(path: Path) -> int | None:
    stem = path.stem
    m = re.search(r"(20\d{2})", stem)
    if m:
        return int(m.group(1))
    m = re.search(
        r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)\s+(\d{4})",
        stem,
        flags=re.I,
    )
    if m:
        return int(m.group(2))
    m = re.search(r"\b(\d{2})\b", stem)
    if m:
        y = int(m.group(1))
        if y < 50:
            return 2000 + y
    return None


def fallback_year(path: Path) -> int | None:
    """When filename has Spanish month without year (e.g. MARZO), infer from neighbouring workbooks."""
    y = year_from_workbook_filename(path)
    if y is not None:
        return y
    parent = path.parent
    if not parent.is_dir():
        return None
    years: list[int] = []
    for p in parent.glob("*.xlsx"):
        if p.name.lower() == "no_show_extractions.xlsx":
            continue
        yy = year_from_workbook_filename(p)
        if yy is not None:
            years.append(yy)
    if not years:
        return 2026
    return max(set(years), key=years.count)


def find_no_show_block(ws: openpyxl.worksheet.worksheet.Worksheet) -> (
    tuple[int, int, int, int | None, int | None] | None
):
    """Return (banner_row, header_row, name_col, unidad_col, coment_col); data begins at header_row+1."""
    max_r = min(ws.max_row or 0, 500)
    max_c = min(ws.max_column or 0, 40)
    banner_row = None
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if looks_like_no_show_banner(v) and row_has_checkout_banner(ws, r, max_c):
                banner_row = r
                break
        if banner_row is not None:
            break
    if banner_row is None:
        return None

    header_row = banner_row + 1
    nombre_cols: list[int] = []
    for c in range(1, max_c + 1):
        if is_nombre_huesped_header(ws.cell(header_row, c).value):
            nombre_cols.append(c)
    if not nombre_cols:
        return None

    j0 = max(nombre_cols)
    headers: dict[int, str] = {}
    for c in range(j0, max_c + 1):
        v = ws.cell(header_row, c).value
        if v is None or not str(v).strip():
            if c > j0 + 8:
                break
            continue
        headers[c] = str(v).strip()

    unidad_col = None
    coment_col = None
    for c in sorted(headers):
        h = headers[c]
        if match_header(h, "UNIDAD") or match_header(h, "HABITACION"):
            unidad_col = c
        if "COMENTARIO" in norm_text(h):
            coment_col = c

    if unidad_col is None:
        if j0 + 2 in headers:
            unidad_col = j0 + 2
        else:
            return None

    # Only use a COMENTARIOS column if the header actually says so (layouts vary).
    if coment_col is None:
        for c in range(j0 + 1, max_c + 1):
            if "COMENTARIO" in norm_text(headers.get(c, "")):
                coment_col = c
                break

    return banner_row, header_row, j0, unidad_col, coment_col


def row_has_no_show_data(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    r: int,
    name_col: int,
    unidad_col: int | None,
    coment_col: int | None,
) -> bool:
    nv = ws.cell(r, name_col).value
    if nv is not None and str(nv).strip():
        return True
    if unidad_col:
        uv = ws.cell(r, unidad_col).value
        if uv is not None and str(uv).strip():
            return True
    if coment_col:
        cv = ws.cell(r, coment_col).value
        if cv is not None and str(cv).strip():
            return True
    return False


def extract_sheet_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    block: tuple[int, int, int, int | None, int | None],
) -> list[tuple[str | None, str | None, str | None]]:
    _, header_row, name_col, unidad_col, coment_col = block
    out: list[tuple[str | None, str | None, str | None]] = []
    r = header_row + 1
    max_r = ws.max_row or 0
    empty_streak = 0
    while r <= max_r and empty_streak < 5:
        if row_has_no_show_data(ws, r, name_col, unidad_col, coment_col):
            empty_streak = 0
            def clean(v: object) -> str | None:
                if v is None:
                    return None
                if isinstance(v, float) and v == int(v):
                    return str(int(v))
                return str(v).strip() or None

            n = clean(ws.cell(r, name_col).value)
            u = clean(ws.cell(r, unidad_col).value) if unidad_col else None
            c = clean(ws.cell(r, coment_col).value) if coment_col else None
            if n or u or c:
                out.append((n, u, c))
        else:
            empty_streak += 1
        r += 1
    return out


def resolve_operation_date(
    wb_path: Path, sheet_name: str, ws: openpyxl.worksheet.worksheet.Worksheet
) -> str:
    sy = fallback_year(wb_path)
    d = parse_sheet_name_date(sheet_name)
    if d:
        return d.strftime("%Y-%m-%d")

    cell2 = ws.cell(2, 2).value
    d2 = parse_spanish_date_row(cell2, sy)
    if d2:
        return d2.strftime("%Y-%m-%d")

    if sy:
        d3 = parse_sheet_name_date(sheet_name[:6]) if len(sheet_name.strip()) >= 6 else None
        if d3:
            return d3.strftime("%Y-%m-%d")
    return sheet_name.strip()


def main() -> None:
    files = sorted(LIBRO_DIR.glob("*.xlsx"))
    files = [f for f in files if f.name.lower() != "no_show_extractions.xlsx"]

    rows_out: list[list[object]] = []
    rows_out.append(
        [
            "Name of workbook",
            "Exact date",
            "NOMBRE DEL HUÉSPED",
            "UNIDAD",
            "COMENTARIOS",
        ]
    )

    for wb_path in files:
        try:
            wb = openpyxl.load_workbook(wb_path, data_only=True)
        except Exception as exc:
            rows_out.append([wb_path.name, f"(error abrir: {exc})", None, None, None])
            continue

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            block = find_no_show_block(ws)
            if not block:
                continue
            *_, name_col, unidad_col, coment_col = block
            extracted = extract_sheet_rows(ws, block)
            if not extracted:
                continue
            date_str = resolve_operation_date(wb_path, sheet_name, ws)
            for n, u, c in extracted:
                rows_out.append([wb_path.name, date_str, n, u, c])
        wb.close()

    total = LIBRO_DIR / "No_Show_Extractions.xlsx"
    out_wb = openpyxl.Workbook()
    ws_out = out_wb.active
    assert ws_out is not None
    ws_out.title = "No Show Extractions"

    # Write without destroying title
    title = ws_out.title
    for rr, row in enumerate(rows_out, start=1):
        for cc, val in enumerate(row, start=1):
            ws_out.cell(rr, cc, val)
    ws_out.title = title

    out_wb.save(total)
    print(f"Saved {total} with {len(rows_out) - 1} data rows (+ header).")


if __name__ == "__main__":
    main()
